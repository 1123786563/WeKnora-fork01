"""Section headers/footers must survive DOCX parsing (#3849).

docx routes to Docx2Parser (FirstParser: markitdown, then the python-docx
DocxParser). Neither engine reads the header/footer parts, so header text
was dropped entirely. Docx2Parser now prepends ``[Page Header]`` /
``[Page Footer]`` marker lines collected from every section.
"""

import io
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from docx import Document as WordDocument
from docx.enum.section import WD_SECTION

from docreader.parser import docx_parser as docx_parser_module
from docreader.parser.docx2_parser import Docx2Parser
from docreader.parser.markitdown_parser import StdMarkitdownParser


def _docx_bytes(build):
    doc = WordDocument()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class _FakeManager:
    """Mirror test_docx_tables: run page workers on threads, plain list."""

    def __enter__(self):
        self._items = []
        return self

    def __exit__(self, *exc):
        return False

    def list(self):
        return self._items


def _parse(content):
    return Docx2Parser(file_name="sample.docx", file_type="docx").parse_into_text(
        content
    )


class DocxHeaderFooterTest(unittest.TestCase):
    def _multi_section_bytes(self):
        def build(doc):
            doc.add_paragraph("Body first paragraph")
            doc.add_paragraph("Body second paragraph")
            first = doc.sections[0]
            first.header.is_linked_to_previous = False
            first.header.paragraphs[0].text = "Confidential Header Alpha"
            first.footer.is_linked_to_previous = False
            first.footer.paragraphs[0].text = "Footer Page 1"
            second = doc.add_section(WD_SECTION.NEW_PAGE)
            second.header.is_linked_to_previous = False
            second.header.paragraphs[0].text = "Section Two Header"
            # second.footer stays linked and inherits "Footer Page 1"
            doc.add_paragraph("Second section body")

        return _docx_bytes(build)

    def test_headers_and_footers_prepend_body_with_markers(self):
        document = _parse(self._multi_section_bytes())
        content = document.content

        self.assertIn("[Page Header] Confidential Header Alpha", content)
        self.assertIn("[Page Header] Section Two Header", content)
        self.assertIn("[Page Footer] Footer Page 1", content)
        # Markers sit at the very start, before any body text.
        self.assertTrue(content.startswith("[Page Header] Confidential Header Alpha"))
        for marker in ("[Page Header]", "[Page Footer]"):
            self.assertLess(content.index(marker), content.index("Body first"))
        # Body order is untouched.
        self.assertLess(
            content.index("Body first paragraph"), content.index("Second section body")
        )

    def test_linked_footer_repeat_is_kept_once(self):
        document = _parse(self._multi_section_bytes())
        self.assertEqual(document.content.count("[Page Footer] Footer Page 1"), 1)

    def test_empty_header_section_is_skipped(self):
        def build(doc):
            doc.add_paragraph("Body text")
            first = doc.sections[0]
            first.header.is_linked_to_previous = False  # header part, no text
            second = doc.add_section(WD_SECTION.NEW_PAGE)
            second.header.is_linked_to_previous = False
            second.header.paragraphs[0].text = "Only Real Header"
            doc.add_paragraph("More body")

        document = _parse(_docx_bytes(build))
        self.assertEqual(document.content.count("[Page Header]"), 1)
        self.assertIn("[Page Header] Only Real Header", document.content)

    def test_repeated_header_across_sections_kept_once(self):
        def build(doc):
            doc.add_paragraph("Body text")
            first = doc.sections[0]
            first.header.is_linked_to_previous = False
            first.header.paragraphs[0].text = "Same Running Header"
            second = doc.add_section(WD_SECTION.NEW_PAGE)
            second.header.is_linked_to_previous = False
            second.header.paragraphs[0].text = "Same Running Header"
            doc.add_paragraph("More body")

        document = _parse(_docx_bytes(build))
        self.assertEqual(document.content.count("[Page Header] Same Running Header"), 1)

    def test_multi_paragraph_header_collapses_to_one_line(self):
        def build(doc):
            doc.add_paragraph("Body text")
            header = doc.sections[0].header
            header.is_linked_to_previous = False
            header.paragraphs[0].text = "Alpha Line"
            header.add_paragraph("Beta Line")

        document = _parse(_docx_bytes(build))
        self.assertIn("[Page Header] Alpha Line Beta Line", document.content)
        self.assertEqual(document.content.count("[Page Header]"), 1)

    def test_document_without_headers_has_no_marker_lines(self):
        def build(doc):
            doc.add_paragraph("Plain body paragraph")

        content = _parse(_docx_bytes(build)).content
        self.assertNotIn("[Page Header]", content)
        self.assertNotIn("[Page Footer]", content)
        self.assertIn("Plain body paragraph", content)

    def test_headers_prepend_on_python_docx_fallback(self):
        """When markitdown fails, the python-docx fallback still gets markers."""
        with (
            patch.object(docx_parser_module, "Manager", _FakeManager),
            patch.object(docx_parser_module, "ProcessPoolExecutor", ThreadPoolExecutor),
            patch.object(
                StdMarkitdownParser,
                "_convert_markitdown",
                side_effect=RuntimeError("markitdown unavailable"),
            ),
        ):
            document = _parse(self._multi_section_bytes())

        self.assertIn("[Page Header] Confidential Header Alpha", document.content)
        self.assertIn("[Page Footer] Footer Page 1", document.content)
        self.assertIn("Body first paragraph", document.content)


if __name__ == "__main__":
    unittest.main()
